"""Read an engineer's workbook without saving it or executing VBA.

Usage: bundled-python scripts/audit-engineer-workbook.py source.xlsm output.json
Cached Excel results are evidence, never presented as a fresh recalculation.
"""
import collections
import hashlib
import json
from pathlib import Path
import re
import sys
import warnings
import zipfile
from xml.etree import ElementTree as ET
import openpyxl

source, destination = map(Path, sys.argv[1:3])
with warnings.catch_warnings():
    warnings.simplefilter('ignore')
    formulas = openpyxl.load_workbook(source, data_only=False, keep_links=True)
    values = openpyxl.load_workbook(source, data_only=True, keep_links=True)

report = {'source': {'file': source.name, 'sha256': hashlib.sha256(source.read_bytes()).hexdigest()},
          'method': 'Read-only OOXML inspection; values are saved Excel caches, VBA not executed.', 'sheets': []}
for ws in formulas:
    cached = values[ws.title]
    cells, functions, errors, comments, item_rows = {}, collections.Counter(), [], [], []
    for row in ws:
        for cell in row:
            if cell.value is None:
                continue
            value = cached[cell.coordinate].value
            formula = cell.value if cell.data_type == 'f' else None
            cells[cell.coordinate] = {'value': value, **({'formula': formula} if formula else {})}
            if formula:
                functions.update(re.findall(r'\b([A-Z][A-Z0-9._]*)\s*\(', formula.upper()))
            if cached[cell.coordinate].data_type == 'e':
                errors.append({'cell': cell.coordinate, 'error': value, 'formula': formula})
            if cell.comment:
                comments.append({'cell': cell.coordinate, 'text': cell.comment.text})
    if ws.title in ['Metrado', 'Acero', 'RESUMEN']:
        for r in range(17, ws.max_row + 1):
            code = str(cached[f'C{r}'].value or '').strip()
            if re.fullmatch(r'(?:O\.E\.|OE\.|HU\.)?\d+(?:\.\d+)*', code):
                item_rows.append({'row': r, 'code': code, 'description': cached[f'D{r}'].value,
                                  'unit': cached[f'E{r}'].value, 'total': cached[f'P{r}'].value if ws.title == 'Metrado' else None})
    report['sheets'].append({'name': ws.title, 'state': ws.sheet_state, 'maxRow': ws.max_row, 'maxCol': ws.max_column,
                            'formulaCount': sum('formula' in c for c in cells.values()), 'functions': dict(functions),
                            'errors': errors, 'comments': comments, 'items': item_rows, 'cells': cells,
                            'merges': [str(m) for m in ws.merged_cells.ranges],
                            'printArea': str(ws.print_area), 'printTitles': ws.print_title_rows})
with zipfile.ZipFile(source) as archive:
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    workbook = ET.fromstring(archive.read('xl/workbook.xml'))
    report['definedNames'] = [dict(name.attrib, value=name.text) for name in workbook.findall('s:definedNames/s:definedName', ns)]
    report['externalLinks'] = [name for name in archive.namelist() if re.fullmatch(r'xl/externalLinks/externalLink\d+.xml', name)]
    report['hasMacros'] = 'xl/vbaProject.bin' in archive.namelist()
destination.parent.mkdir(parents=True, exist_ok=True)
destination.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf8')
print(json.dumps({**report['source'], 'sheets': [{k: s[k] for k in ['name', 'state', 'maxRow', 'formulaCount', 'functions', 'errors']} for s in report['sheets']]}, ensure_ascii=False, default=str))
